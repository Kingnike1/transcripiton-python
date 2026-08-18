"""Meeting export service for portable and document formats."""

from __future__ import annotations

import io
import json
import re
from dataclasses import asdict, dataclass
from html import escape
from typing import Any, Callable

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy.orm import Session

from app.models.analysis import MeetingAnalysis
from app.models.audio import Audio
from app.models.meeting import Meeting
from app.models.participant import Participant
from app.models.transcription import Transcription


@dataclass(frozen=True)
class ExportParticipant:
    speaker_label: str
    display_name: str | None
    confirmed: bool


@dataclass(frozen=True)
class ExportSegment:
    start_time: float
    end_time: float
    speaker_label: str | None
    speaker_name: str | None
    text: str


@dataclass(frozen=True)
class ExportBundle:
    meeting_id: int
    title: str
    description: str | None
    status: str
    created_at: str
    participants: list[ExportParticipant]
    transcript: str | None
    segments: list[ExportSegment]
    analysis: dict[str, Any] | None


@dataclass(frozen=True)
class ExportDocument:
    content: bytes
    media_type: str
    filename: str


class ExportService:
    """Build and render an ownership-aware meeting export."""

    FORMATS = {"txt", "md", "json", "docx", "pdf"}

    def __init__(self, db: Session) -> None:
        self.db = db

    def build_bundle(self, meeting_id: int, owner_id: int | None = None) -> ExportBundle:
        query = self.db.query(Meeting).filter(Meeting.id == meeting_id, Meeting.deleted_at.is_(None))
        if owner_id is not None:
            query = query.filter(Meeting.owner_id == owner_id)
        meeting = query.one_or_none()
        if meeting is None:
            raise ValueError("Meeting not found")

        participants = (
            self.db.query(Participant)
            .filter(Participant.meeting_id == meeting.id)
            .order_by(Participant.speaker_label.asc())
            .all()
        )
        participant_map = {
            item.speaker_label: item.display_name
            if item.confirmed and item.display_name
            else item.speaker_label
            for item in participants
        }

        audio = (
            self.db.query(Audio)
            .filter(Audio.meeting_id == meeting.id, Audio.deleted_at.is_(None))
            .order_by(Audio.created_at.desc())
            .first()
        )
        transcription: Transcription | None = None
        if audio is not None:
            transcription = (
                self.db.query(Transcription)
                .filter(Transcription.audio_id == audio.id, Transcription.deleted_at.is_(None))
                .one_or_none()
            )

        segments: list[ExportSegment] = []
        if transcription is not None:
            if transcription.speaker_segments:
                for speaker_segment in transcription.speaker_segments:
                    label = speaker_segment.speaker_label
                    segments.append(
                        ExportSegment(
                            start_time=speaker_segment.start_time,
                            end_time=speaker_segment.end_time,
                            speaker_label=label,
                            speaker_name=participant_map.get(label, label),
                            text=speaker_segment.text,
                        )
                    )
            else:
                for transcript_segment in transcription.segments:
                    segments.append(
                        ExportSegment(
                            start_time=transcript_segment.start_time,
                            end_time=transcript_segment.end_time,
                            speaker_label=None,
                            speaker_name=None,
                            text=transcript_segment.text,
                        )
                    )

        analysis_row = self.db.query(MeetingAnalysis).filter_by(meeting_id=meeting.id).one_or_none()
        analysis = self._analysis_payload(analysis_row) if analysis_row is not None else None
        return ExportBundle(
            meeting_id=meeting.id,
            title=meeting.title,
            description=meeting.description,
            status=meeting.status,
            created_at=meeting.created_at.isoformat(),
            participants=[
                ExportParticipant(
                    speaker_label=item.speaker_label,
                    display_name=item.display_name,
                    confirmed=item.confirmed,
                )
                for item in participants
            ],
            transcript=transcription.text if transcription is not None else None,
            segments=segments,
            analysis=analysis,
        )

    def export(self, meeting_id: int, fmt: str, owner_id: int | None = None) -> ExportDocument:
        normalized = fmt.lower().strip()
        if normalized not in self.FORMATS:
            raise ValueError(f"Unsupported export format: {fmt}")
        bundle = self.build_bundle(meeting_id, owner_id=owner_id)
        slug = self._slug(bundle.title) or f"meeting-{bundle.meeting_id}"
        renderers: dict[str, tuple[Callable[[ExportBundle], bytes], str]] = {
            "txt": (self._render_txt, "text/plain; charset=utf-8"),
            "md": (self._render_markdown, "text/markdown; charset=utf-8"),
            "json": (self._render_json, "application/json"),
            "docx": (
                self._render_docx,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            "pdf": (self._render_pdf, "application/pdf"),
        }
        renderer, media_type = renderers[normalized]
        return ExportDocument(
            content=renderer(bundle),
            media_type=media_type,
            filename=f"{slug}.{normalized}",
        )

    @staticmethod
    def _analysis_payload(row: MeetingAnalysis) -> dict[str, Any]:
        def load_items(value: str | None) -> list[Any]:
            if not value:
                return []
            try:
                loaded = json.loads(value)
                return loaded if isinstance(loaded, list) else []
            except json.JSONDecodeError:
                return []

        return {
            "summary": row.summary,
            "action_items": load_items(row.action_items),
            "decisions": load_items(row.decisions),
            "risks": load_items(row.risks),
            "open_questions": load_items(row.open_questions),
            "follow_up_tasks": load_items(row.follow_up_tasks),
            "provider": row.provider,
            "model_name": row.model_name,
        }

    @staticmethod
    def _slug(value: str) -> str:
        normalized = re.sub(r"[^a-zA-Z0-9_-]+", "-", value.strip()).strip("-")
        return normalized[:80].lower()

    @staticmethod
    def _format_time(seconds: float) -> str:
        total = max(0, int(seconds))
        return f"{total // 60:02d}:{total % 60:02d}"

    def _transcript_lines(self, bundle: ExportBundle) -> list[str]:
        if bundle.segments:
            lines: list[str] = []
            for segment in bundle.segments:
                who = segment.speaker_name or segment.speaker_label or "Falante"
                lines.append(
                    f"[{self._format_time(segment.start_time)}-{self._format_time(segment.end_time)}] "
                    f"{who}: {segment.text}"
                )
            return lines
        return [bundle.transcript] if bundle.transcript else ["Transcrição ainda não disponível."]

    @staticmethod
    def _item_text(item: Any) -> str:
        if isinstance(item, dict):
            primary = item.get("text") or item.get("title") or item.get("task") or item.get("description")
            if primary:
                evidence = item.get("evidence")
                return f"{primary} — evidência: {evidence}" if evidence else str(primary)
            return json.dumps(item, ensure_ascii=False)
        return str(item)

    def _analysis_lines(self, bundle: ExportBundle) -> list[str]:
        if bundle.analysis is None:
            return ["Análise ainda não disponível."]
        lines = [f"Resumo: {bundle.analysis.get('summary') or 'Não disponível.'}"]
        sections = [
            ("Ações", "action_items"),
            ("Decisões", "decisions"),
            ("Riscos", "risks"),
            ("Perguntas em aberto", "open_questions"),
            ("Follow-ups", "follow_up_tasks"),
        ]
        for title, key in sections:
            lines.append(f"\n{title}:")
            items = bundle.analysis.get(key) or []
            lines.extend([f"- {self._item_text(item)}" for item in items] or ["- Nenhum item."])
        return lines

    def _render_txt(self, bundle: ExportBundle) -> bytes:
        lines = [
            bundle.title,
            "=" * len(bundle.title),
            f"Status: {bundle.status}",
            f"Criada em: {bundle.created_at}",
            f"Descrição: {bundle.description or 'Sem descrição.'}",
            "",
            "PARTICIPANTES",
        ]
        lines.extend(
            [
                f"- {p.display_name or p.speaker_label} ({p.speaker_label})"
                f"{' ✓' if p.confirmed else ''}"
                for p in bundle.participants
            ]
            or ["- Nenhum participante identificado."]
        )
        lines.extend(
            [
                "",
                "TRANSCRIÇÃO",
                *self._transcript_lines(bundle),
                "",
                "ANÁLISE",
                *self._analysis_lines(bundle),
            ]
        )
        return "\n".join(lines).encode("utf-8")

    def _render_markdown(self, bundle: ExportBundle) -> bytes:
        lines = [
            f"# {bundle.title}",
            "",
            f"**Status:** {bundle.status}  ",
            f"**Criada em:** {bundle.created_at}  ",
            f"**Descrição:** {bundle.description or 'Sem descrição.'}",
            "",
            "## Participantes",
        ]
        lines.extend(
            [
                f"- {p.display_name or p.speaker_label} (`{p.speaker_label}`)"
                f"{' ✓' if p.confirmed else ''}"
                for p in bundle.participants
            ]
            or ["- Nenhum participante identificado."]
        )
        lines.extend(["", "## Transcrição"])
        lines.extend([f"- {line}" for line in self._transcript_lines(bundle)])
        lines.extend(["", "## Análise"])
        lines.extend(self._analysis_lines(bundle))
        return "\n".join(lines).encode("utf-8")

    @staticmethod
    def _render_json(bundle: ExportBundle) -> bytes:
        return json.dumps(asdict(bundle), ensure_ascii=False, indent=2).encode("utf-8")

    def _render_docx(self, bundle: ExportBundle) -> bytes:
        document = Document()
        document.add_heading(bundle.title, level=0)
        document.add_paragraph(f"Status: {bundle.status}")
        document.add_paragraph(f"Criada em: {bundle.created_at}")
        document.add_paragraph(bundle.description or "Sem descrição.")
        document.add_heading("Participantes", level=1)
        if bundle.participants:
            for participant in bundle.participants:
                document.add_paragraph(
                    f"{participant.display_name or participant.speaker_label} ({participant.speaker_label})"
                    f"{' — confirmado' if participant.confirmed else ''}",
                    style="List Bullet",
                )
        else:
            document.add_paragraph("Nenhum participante identificado.")
        document.add_heading("Transcrição", level=1)
        for line in self._transcript_lines(bundle):
            document.add_paragraph(line)
        document.add_heading("Análise", level=1)
        for line in self._analysis_lines(bundle):
            document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _pdf_safe(value: str) -> str:
        windows_text = value.encode("cp1252", errors="replace").decode("cp1252")
        return escape(windows_text)

    def _render_pdf(self, bundle: ExportBundle) -> bytes:
        buffer = io.BytesIO()
        styles = getSampleStyleSheet()
        story = [Paragraph(self._pdf_safe(bundle.title), styles["Title"]), Spacer(1, 12)]
        for line in [
            f"Status: {bundle.status}",
            f"Criada em: {bundle.created_at}",
            bundle.description or "Sem descrição.",
        ]:
            story.append(Paragraph(self._pdf_safe(line), styles["BodyText"]))
        story.extend([Spacer(1, 12), Paragraph("Participantes", styles["Heading2"])])
        participant_lines = [
            f"- {p.display_name or p.speaker_label} ({p.speaker_label})"
            f"{' — confirmado' if p.confirmed else ''}"
            for p in bundle.participants
        ] or ["Nenhum participante identificado."]
        for line in participant_lines:
            story.append(Paragraph(self._pdf_safe(line), styles["BodyText"]))
        story.extend([Spacer(1, 12), Paragraph("Transcrição", styles["Heading2"])])
        for line in self._transcript_lines(bundle):
            story.append(Paragraph(self._pdf_safe(line), styles["BodyText"]))
        story.extend([Spacer(1, 12), Paragraph("Análise", styles["Heading2"])])
        for line in self._analysis_lines(bundle):
            story.append(Paragraph(self._pdf_safe(line), styles["BodyText"]))
        SimpleDocTemplate(buffer, pagesize=A4, title=self._pdf_safe(bundle.title)).build(story)
        return buffer.getvalue()
